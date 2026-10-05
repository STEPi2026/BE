"""Seed STEPi reference data from the functional-specification workbook.

Usage:
    python seed.py
    python seed.py --workbook "/mnt/data/STEPi 기능명세서.xlsx"

The script deliberately does not create tables. Apply Alembic migrations first.
"""

from __future__ import annotations

import argparse
import sys
from collections.abc import Mapping
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.db.database import SessionLocal
from app.db.models.badge import Badge
from app.db.models.concept import Concept
from app.db.models.misconception import Misconception
from app.db.models.problem import Problem
from app.db.models.problem_concept import ProblemConcept
from app.db.models.problem_skill import ProblemSkill
from app.db.models.skill import Skill
from app.db.models.skill_prerequisite import SkillPrerequisite


# ============================================================
# 기본 설정
# ============================================================

DEFAULT_WORKBOOK = Path(
    "/Users/gnswl/Downloads/STEPi 기능명세서.xlsx"
)


# ============================================================
# 시트 탐색용 키워드
# ============================================================

SHEET_KEYWORDS = {
    "concepts": ("Concept",),
    "skills": ("Skill",),
    "prerequisites": ("Prerequisite",),
    "misconceptions": ("Misconception", "오개념"),
    "problems": ("Problem",),
    "problem_concepts": (
        "Problem-Concept",
        "Problem Concept",
    ),
    "problem_skills": (
        "Problem-Skill",
        "Problem Skill",
    ),
    "badges": ("Badge",),
}


# ============================================================
# Excel 헤더 매핑
# ============================================================

HEADERS = {
    "concepts": {
        "code": (
            "conceptcode",
            "코드",
        ),
        "name": (
            "concept명",
            "conceptname",
            "개념명",
            "name",
        ),
        "description": (
            "설명",
            "description",
        ),
        "subject": (
            "대단원",
            "과목",
            "subject",
        ),
        "grade_level": (
            "학년과정",
            "학년",
            "gradelevel",
            "grade",
        ),
        "parent_code": (
            "parentconceptcode",
            "상위conceptcode",
            "상위개념코드",
        ),
    },

    "skills": {
        "code": (
            "skillcode",
            "코드",
        ),
        "name": (
            "skill명",
            "skillname",
            "스킬명",
            "name",
        ),
        "description": (
            "설명",
            "description",
        ),
        "concept_code": (
            "conceptcode",
            "개념코드",
        ),
    },

    "prerequisites": {
        "skill_code": (
            "skillcode",
        ),
        "prerequisite_skill_code": (
            "prerequisiteskillcode",
            "선수skillcode",
            "선수스킬코드",
        ),
        "description": (
            "관계설명",
            "설명",
            "description",
        ),
    },

    "misconceptions": {
        "code": (
            "misconceptioncode",
            "오개념코드",
            "코드",
        ),
        "name": (
            "misconception명",
            "misconceptionname",
            "오개념명",
            "name",
        ),
        "description": (
            "설명",
            "description",
        ),
        "concept_code": (
            "conceptcode",
            "개념코드",
        ),
    },

    "problems": {
        "problem_code": (
            "problemcode",
            "문제코드",
        ),
        "question": (
            "문제",
            "문제내용",
            "question",
        ),
        "answer": (
            "정답",
            "answer",
        ),
        "difficulty": (
            "난이도",
            "difficulty",
        ),
        "problem_type": (
            "문제유형",
            "problemtype",
            "유형",
        ),
        "usage": (
            "용도",
            "usage",
            "사용처",
        ),
        "explanation": (
            "해설",
            "explanation",
        ),
    },

    "problem_concepts": {
        "problem_code": (
            "problemcode",
            "문제코드",
        ),
        "concept_code": (
            "conceptcode",
            "개념코드",
        ),
    },

    "problem_skills": {
        "problem_code": (
            "problemcode",
            "문제코드",
        ),
        "skill_code": (
            "skillcode",
            "스킬코드",
        ),
    },

    "badges": {
        "code": (
            "badgecode",
            "배지코드",
            "코드",
        ),
        "name": (
            "badge명",
            "badgename",
            "배지명",
            "name",
        ),
        "description": (
            "설명",
            "description",
        ),
        "condition_type": (
            "획득조건유형",
            "conditiontype",
            "조건유형",
        ),
        "threshold": (
            "threshold",
            "기준값",
            "달성기준",
        ),
        "reward_xp": (
            "rewardxp",
            "보상xp",
            "xp보상",
        ),
    },
}


# ============================================================
# 현재 기준 데이터 예상 개수
#
# Misconception은 개수를 고정하지 않는다.
# 나중에 Excel에 오개념을 추가하면 자동 반영된다.
# ============================================================

EXPECTED_COUNTS = {
    "concepts": 42,
    "skills": 94,
    "prerequisites": 92,
    "problems": 94,
    "problem_concepts": 94,
    "problem_skills": 94,
    "badges": 6,
}


class SeedValidationError(ValueError):
    """Seed 원본 데이터 검증 오류."""


# ============================================================
# 공통 유틸
# ============================================================

def normalized(value: Any) -> str:
    return "".join(
        char
        for char in str(value or "").lower()
        if char.isalnum()
    )


def text(value: Any) -> str | None:
    if value is None:
        return None

    value = str(value).strip()

    return value or None


def required(
    row: Mapping[str, Any],
    field: str,
    row_number: int,
    dataset: str,
) -> Any:

    value = row.get(field)

    if value is None or (
        isinstance(value, str)
        and not value.strip()
    ):
        raise SeedValidationError(
            f"{dataset} sheet row {row_number}: "
            f"missing required '{field}'"
        )

    return value


def integer(
    value: Any,
    field: str,
    row_number: int,
    dataset: str,
) -> int:

    try:
        raw = required(
            {field: value},
            field,
            row_number,
            dataset,
        )

        number = float(raw)

    except (TypeError, ValueError) as error:
        raise SeedValidationError(
            f"{dataset} sheet row {row_number}: "
            f"invalid integer '{field}'"
        ) from error

    if not number.is_integer():
        raise SeedValidationError(
            f"{dataset} sheet row {row_number}: "
            f"'{field}' must be an integer"
        )

    return int(number)


# ============================================================
# 시트 찾기
# ============================================================

def find_sheet(
    workbook: Any,
    dataset: str,
) -> Any:

    keywords = SHEET_KEYWORDS[dataset]

    matches = [
        ws
        for ws in workbook.worksheets
        if any(
            keyword.lower() in ws.title.lower()
            for keyword in keywords
        )
    ]

    if not matches:
        raise SeedValidationError(
            f"missing worksheet for {dataset}; "
            f"looked for {keywords}"
        )

    # Problem은 Problem-Concept / Problem-Skill과
    # 이름이 겹칠 수 있으므로 관계 시트 제외
    if dataset == "problems":

        filtered = [
            ws
            for ws in matches
            if "concept" not in ws.title.lower()
            and "skill" not in ws.title.lower()
        ]

        if filtered:
            matches = filtered

    # Skill 역시 Skill-Prerequisite / Problem-Skill과
    # 이름이 겹칠 수 있으므로 제외
    if dataset == "skills":

        filtered = [
            ws
            for ws in matches
            if "prerequisite" not in ws.title.lower()
            and "problem" not in ws.title.lower()
        ]

        if filtered:
            matches = filtered

    return matches[0]


# ============================================================
# 헤더 행 찾기
#
# 중요:
# read_only=True에서는 worksheet.max_row가 None일 수 있으므로
# max_row를 사용하지 않는다.
# ============================================================

def find_header_row(
    worksheet: Any,
    dataset: str,
) -> tuple[int, dict[str, int]]:

    aliases = HEADERS[dataset]

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            values_only=True,
        ),
        start=1,
    ):

        # 상단 30행까지만 헤더 탐색
        if row_number > 30:
            break

        normalized_cells = {
            normalized(value): index
            for index, value in enumerate(row)
            if value is not None
        }

        mapping: dict[str, int] = {}

        for field, candidates in aliases.items():

            for candidate in candidates:

                key = normalized(candidate)

                if key in normalized_cells:
                    mapping[field] = normalized_cells[key]
                    break

        # 최소 2개의 헤더가 일치하면 후보로 인정
        if len(mapping) >= 2:
            return row_number, mapping

    raise SeedValidationError(
        f"could not locate header row "
        f"for dataset '{dataset}' "
        f"in worksheet '{worksheet.title}'"
    )


# ============================================================
# Excel 데이터 읽기
# ============================================================

def read_rows(
    workbook: Any,
    dataset: str,
) -> list[tuple[int, dict[str, Any]]]:

    worksheet = find_sheet(
        workbook,
        dataset,
    )

    header_row, mapping = find_header_row(
        worksheet,
        dataset,
    )

    rows: list[
        tuple[int, dict[str, Any]]
    ] = []

    for row_number, values in enumerate(
        worksheet.iter_rows(
            min_row=header_row + 1,
            values_only=True,
        ),
        start=header_row + 1,
    ):

        row: dict[str, Any] = {}

        for field, column_index in mapping.items():

            if column_index < len(values):
                row[field] = values[column_index]
            else:
                row[field] = None

        # 완전히 빈 행 제외
        if not any(
            text(value)
            for value in row.values()
        ):
            continue

        # 반복 헤더 제외
        normalized_values = {
            normalized(value)
            for value in row.values()
            if value is not None
        }

        header_aliases = {
            normalized(alias)
            for aliases in HEADERS[dataset].values()
            for alias in aliases
        }

        if (
            normalized_values
            and normalized_values <= header_aliases
        ):
            continue

        rows.append(
            (
                row_number,
                row,
            )
        )

    return rows


# ============================================================
# 중복 검사
# ============================================================

def ensure_unique(
    rows: list[tuple[int, dict[str, Any]]],
    field: str,
    dataset: str,
) -> None:

    seen: dict[str, int] = {}

    for row_number, row in rows:

        value = str(
            required(
                row,
                field,
                row_number,
                dataset,
            )
        ).strip()

        if value in seen:

            raise SeedValidationError(
                f"{dataset} sheet row {row_number}: "
                f"duplicate '{value}' "
                f"(first seen row {seen[value]})"
            )

        seen[value] = row_number


# ============================================================
# Skill-Prerequisite 정제
#
# Excel 안의
# SKL-2 방정식과 부등식
# SKL-3 경우의 수
# SKL-4 행렬
#
# 같은 단원 구분행을 자동 제외한다.
# ============================================================

def filter_prerequisites(
    data: dict[
        str,
        list[tuple[int, dict[str, Any]]],
    ],
) -> None:

    skill_codes = {
        str(row["code"]).strip()
        for _, row in data["skills"]
        if text(row.get("code"))
    }

    filtered: list[
        tuple[int, dict[str, Any]]
    ] = []

    skipped: list[
        tuple[int, str | None, str | None]
    ] = []

    for row_number, row in data["prerequisites"]:

        skill_code = text(
            row.get("skill_code")
        )

        prerequisite_code = text(
            row.get(
                "prerequisite_skill_code"
            )
        )

        # 실제 존재하는 Skill → Skill 관계만 사용
        if (
            skill_code in skill_codes
            and prerequisite_code in skill_codes
        ):
            filtered.append(
                (
                    row_number,
                    row,
                )
            )

        else:
            skipped.append(
                (
                    row_number,
                    skill_code,
                    prerequisite_code,
                )
            )

    data["prerequisites"] = filtered

    if skipped:

        print(
            "\n[INFO] "
            "Skill-Prerequisite 비데이터 행 제외:"
        )

        for (
            row_number,
            skill_code,
            prerequisite_code,
        ) in skipped:

            print(
                f"  row {row_number}: "
                f"{skill_code!r} -> "
                f"{prerequisite_code!r}"
            )


# ============================================================
# 전체 무결성 검증
# ============================================================

def validate(
    data: dict[
        str,
        list[tuple[int, dict[str, Any]]],
    ],
) -> None:

    # --------------------------------------------------------
    # 데이터 건수
    # --------------------------------------------------------

    for dataset, expected in EXPECTED_COUNTS.items():

        actual = len(
            data[dataset]
        )

        if actual != expected:

            raise SeedValidationError(
                f"{dataset}: "
                f"expected {expected} data rows, "
                f"found {actual}"
            )

    # Misconception은 앞으로 증가할 예정이므로
    # 특정 개수로 고정하지 않는다.
    if not data["misconceptions"]:

        raise SeedValidationError(
            "misconceptions: no data rows found"
        )

    # --------------------------------------------------------
    # 코드 중복
    # --------------------------------------------------------

    for dataset, field in (
        ("concepts", "code"),
        ("skills", "code"),
        ("misconceptions", "code"),
        ("problems", "problem_code"),
        ("badges", "code"),
    ):

        ensure_unique(
            data[dataset],
            field,
            dataset,
        )

    concept_codes = {
        str(row["code"]).strip()
        for _, row in data["concepts"]
    }

    skill_codes = {
        str(row["code"]).strip()
        for _, row in data["skills"]
    }

    problem_codes = {
        str(row["problem_code"]).strip()
        for _, row in data["problems"]
    }

    # --------------------------------------------------------
    # Concept parent
    # --------------------------------------------------------

    for row_number, row in data["concepts"]:

        parent = text(
            row.get("parent_code")
        )

        if (
            parent
            and parent not in concept_codes
        ):

            raise SeedValidationError(
                f"concepts sheet row {row_number}: "
                f"unknown parent concept '{parent}'"
            )

    # --------------------------------------------------------
    # Skill → Concept
    # Misconception → Concept
    # --------------------------------------------------------

    for dataset in (
        "skills",
        "misconceptions",
    ):

        for row_number, row in data[dataset]:

            concept_code = str(
                required(
                    row,
                    "concept_code",
                    row_number,
                    dataset,
                )
            ).strip()

            if concept_code not in concept_codes:

                raise SeedValidationError(
                    f"{dataset} sheet row "
                    f"{row_number}: "
                    f"unknown concept "
                    f"'{concept_code}'"
                )

    # --------------------------------------------------------
    # Skill-Prerequisite
    # --------------------------------------------------------

    prerequisite_pairs: set[
        tuple[str, str]
    ] = set()

    for row_number, row in data["prerequisites"]:

        skill_code = str(
            required(
                row,
                "skill_code",
                row_number,
                "prerequisites",
            )
        ).strip()

        prerequisite_code = str(
            required(
                row,
                "prerequisite_skill_code",
                row_number,
                "prerequisites",
            )
        ).strip()

        if skill_code not in skill_codes:

            raise SeedValidationError(
                f"prerequisites sheet row "
                f"{row_number}: "
                f"unknown skill "
                f"'{skill_code}'"
            )

        if prerequisite_code not in skill_codes:

            raise SeedValidationError(
                f"prerequisites sheet row "
                f"{row_number}: "
                f"unknown prerequisite skill "
                f"'{prerequisite_code}'"
            )

        if skill_code == prerequisite_code:

            raise SeedValidationError(
                f"prerequisites sheet row "
                f"{row_number}: "
                f"skill cannot depend on itself"
            )

        pair = (
            skill_code,
            prerequisite_code,
        )

        if pair in prerequisite_pairs:

            raise SeedValidationError(
                f"prerequisites sheet row "
                f"{row_number}: "
                f"duplicate relationship "
                f"{skill_code} -> "
                f"{prerequisite_code}"
            )

        prerequisite_pairs.add(
            pair
        )

    # --------------------------------------------------------
    # Problem-Concept / Problem-Skill
    # --------------------------------------------------------

    relation_checks = (
        (
            "problem_concepts",
            "concept_code",
            concept_codes,
        ),
        (
            "problem_skills",
            "skill_code",
            skill_codes,
        ),
    )

    for (
        dataset,
        reference_field,
        valid_codes,
    ) in relation_checks:

        seen: set[
            tuple[str, str]
        ] = set()

        for row_number, row in data[dataset]:

            problem_code = str(
                required(
                    row,
                    "problem_code",
                    row_number,
                    dataset,
                )
            ).strip()

            reference_code = str(
                required(
                    row,
                    reference_field,
                    row_number,
                    dataset,
                )
            ).strip()

            if problem_code not in problem_codes:

                raise SeedValidationError(
                    f"{dataset} sheet row "
                    f"{row_number}: "
                    f"unknown problem "
                    f"'{problem_code}'"
                )

            if reference_code not in valid_codes:

                raise SeedValidationError(
                    f"{dataset} sheet row "
                    f"{row_number}: "
                    f"unknown reference "
                    f"'{reference_code}'"
                )

            pair = (
                problem_code,
                reference_code,
            )

            if pair in seen:

                raise SeedValidationError(
                    f"{dataset} sheet row "
                    f"{row_number}: "
                    f"duplicate relationship "
                    f"{pair}"
                )

            seen.add(
                pair
            )

    # --------------------------------------------------------
    # 숫자 데이터
    # --------------------------------------------------------

    for row_number, row in data["problems"]:

        integer(
            row["difficulty"],
            "difficulty",
            row_number,
            "problems",
        )

    for row_number, row in data["badges"]:

        integer(
            row["threshold"],
            "threshold",
            row_number,
            "badges",
        )

        integer(
            row["reward_xp"],
            "reward_xp",
            row_number,
            "badges",
        )


# ============================================================
# UPSERT
# ============================================================

def upsert(
    session: Session,
    model: Any,
    lookup: Mapping[str, Any],
    values: Mapping[str, Any],
) -> Any:

    instance = session.scalar(
        select(model).filter_by(
            **lookup
        )
    )

    if instance is None:

        instance = model(
            **dict(lookup),
            **dict(values),
        )

        session.add(
            instance
        )

    else:

        for key, value in values.items():

            setattr(
                instance,
                key,
                value,
            )

    return instance


# ============================================================
# Seed 실행
# ============================================================

def seed(
    workbook_path: Path,
) -> None:

    if not workbook_path.is_file():

        raise FileNotFoundError(
            f"Workbook not found: "
            f"{workbook_path}"
        )

    print(
        f"[INFO] Workbook: "
        f"{workbook_path}"
    )

    # --------------------------------------------------------
    # Excel 로딩
    # --------------------------------------------------------

    workbook = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    try:

        data = {
            dataset: read_rows(
                workbook,
                dataset,
            )
            for dataset in HEADERS
        }

        # 선수관계 시트 구분행 제거
        filter_prerequisites(
            data
        )

        # ----------------------------------------------------
        # 원본 데이터 건수 출력
        # ----------------------------------------------------

        print(
            "\n=== WORKBOOK DATA ==="
        )

        for dataset in (
            "concepts",
            "skills",
            "prerequisites",
            "misconceptions",
            "problems",
            "problem_concepts",
            "problem_skills",
            "badges",
        ):

            print(
                f"{dataset}: "
                f"{len(data[dataset])}"
            )

        # ----------------------------------------------------
        # Excel 전체 무결성 검증
        # ----------------------------------------------------

        validate(
            data
        )

        print(
            "\n[PASS] "
            "Workbook validation completed."
        )

    finally:

        workbook.close()

    # ========================================================
    # DB 작업 시작
    # ========================================================

    session = SessionLocal()

    try:

        with session.begin():

            # =================================================
            # 1. Concept
            # =================================================

            for _, row in data["concepts"]:

                upsert(
                    session,
                    Concept,
                    {
                        "code": text(
                            row["code"]
                        )
                    },
                    {
                        "name": text(
                            row["name"]
                        ),
                        "description": text(
                            row.get(
                                "description"
                            )
                        ),
                        "subject": text(
                            row["subject"]
                        ),
                        "grade_level": text(
                            row[
                                "grade_level"
                            ]
                        ),
                    },
                )

            session.flush()

            concepts = {
                code: id_
                for code, id_
                in session.execute(
                    select(
                        Concept.code,
                        Concept.id,
                    )
                )
            }

            # =================================================
            # Concept parent 연결
            # =================================================

            for _, row in data["concepts"]:

                parent_code = text(
                    row.get(
                        "parent_code"
                    )
                )

                if not parent_code:
                    continue

                concept = session.scalar(
                    select(
                        Concept
                    ).where(
                        Concept.code
                        == text(
                            row["code"]
                        )
                    )
                )

                if concept is not None:

                    concept.parent_id = (
                        concepts[
                            parent_code
                        ]
                    )

            # =================================================
            # 2. Skill
            # =================================================

            for _, row in data["skills"]:

                concept_code = text(
                    row["concept_code"]
                )

                upsert(
                    session,
                    Skill,
                    {
                        "code": text(
                            row["code"]
                        )
                    },
                    {
                        "name": text(
                            row["name"]
                        ),
                        "description": text(
                            row.get(
                                "description"
                            )
                        ),
                        "concept_id":
                            concepts[
                                concept_code
                            ],
                    },
                )

            # =================================================
            # 3. Misconception
            # =================================================

            for _, row in data["misconceptions"]:

                concept_code = text(
                    row["concept_code"]
                )

                upsert(
                    session,
                    Misconception,
                    {
                        "code": text(
                            row["code"]
                        )
                    },
                    {
                        "name": text(
                            row["name"]
                        ),
                        "description": text(
                            row.get(
                                "description"
                            )
                        ),
                        "concept_id":
                            concepts[
                                concept_code
                            ],
                    },
                )

            # =================================================
            # 4. Problem
            # =================================================

            for row_number, row in data["problems"]:

                upsert(
                    session,
                    Problem,
                    {
                        "problem_code": text(
                            row[
                                "problem_code"
                            ]
                        )
                    },
                    {
                        "question": text(
                            row["question"]
                        ),
                        "answer": text(
                            row["answer"]
                        ),
                        "difficulty": integer(
                            row[
                                "difficulty"
                            ],
                            "difficulty",
                            row_number,
                            "problems",
                        ),
                        "problem_type": text(
                            row[
                                "problem_type"
                            ]
                        ),
                        "usage": text(
                            row["usage"]
                        ),
                        "explanation": text(
                            row.get(
                                "explanation"
                            )
                        ),
                    },
                )

            # =================================================
            # 5. Badge
            # =================================================

            for row_number, row in data["badges"]:

                upsert(
                    session,
                    Badge,
                    {
                        "code": text(
                            row["code"]
                        )
                    },
                    {
                        "name": text(
                            row["name"]
                        ),
                        "description": text(
                            row.get(
                                "description"
                            )
                        ),
                        "condition_type": text(
                            row[
                                "condition_type"
                            ]
                        ),
                        "threshold": integer(
                            row[
                                "threshold"
                            ],
                            "threshold",
                            row_number,
                            "badges",
                        ),
                        "reward_xp": integer(
                            row[
                                "reward_xp"
                            ],
                            "reward_xp",
                            row_number,
                            "badges",
                        ),
                    },
                )

            session.flush()

            # =================================================
            # ID Map
            # =================================================

            skills = {
                code: id_
                for code, id_
                in session.execute(
                    select(
                        Skill.code,
                        Skill.id,
                    )
                )
            }

            problems = {
                code: id_
                for code, id_
                in session.execute(
                    select(
                        Problem.problem_code,
                        Problem.id,
                    )
                )
            }

            # =================================================
            # 6. Skill-Prerequisite
            # =================================================

            for _, row in data["prerequisites"]:

                skill_code = text(
                    row[
                        "skill_code"
                    ]
                )

                prerequisite_code = text(
                    row[
                        "prerequisite_skill_code"
                    ]
                )

                upsert(
                    session,
                    SkillPrerequisite,
                    {
                        "skill_id":
                            skills[
                                skill_code
                            ],

                        "prerequisite_skill_id":
                            skills[
                                prerequisite_code
                            ],
                    },
                    {
                        "description": text(
                            row.get(
                                "description"
                            )
                        )
                    },
                )

            # =================================================
            # 7. Problem-Concept
            # =================================================

            for _, row in data[
                "problem_concepts"
            ]:

                problem_code = text(
                    row[
                        "problem_code"
                    ]
                )

                concept_code = text(
                    row[
                        "concept_code"
                    ]
                )

                upsert(
                    session,
                    ProblemConcept,
                    {
                        "problem_id":
                            problems[
                                problem_code
                            ],

                        "concept_id":
                            concepts[
                                concept_code
                            ],
                    },
                    {},
                )

            # =================================================
            # 8. Problem-Skill
            # =================================================

            for _, row in data[
                "problem_skills"
            ]:

                problem_code = text(
                    row[
                        "problem_code"
                    ]
                )

                skill_code = text(
                    row[
                        "skill_code"
                    ]
                )

                upsert(
                    session,
                    ProblemSkill,
                    {
                        "problem_id":
                            problems[
                                problem_code
                            ],

                        "skill_id":
                            skills[
                                skill_code
                            ],
                    },
                    {},
                )

        # ====================================================
        # 실제 DB 건수 확인
        # ====================================================

        counts = {
            "Concept":
                session.scalar(
                    select(
                        func.count()
                    ).select_from(
                        Concept
                    )
                ),

            "Skill":
                session.scalar(
                    select(
                        func.count()
                    ).select_from(
                        Skill
                    )
                ),

            "SkillPrerequisite":
                session.scalar(
                    select(
                        func.count()
                    ).select_from(
                        SkillPrerequisite
                    )
                ),

            "Misconception":
                session.scalar(
                    select(
                        func.count()
                    ).select_from(
                        Misconception
                    )
                ),

            "Problem":
                session.scalar(
                    select(
                        func.count()
                    ).select_from(
                        Problem
                    )
                ),

            "ProblemConcept":
                session.scalar(
                    select(
                        func.count()
                    ).select_from(
                        ProblemConcept
                    )
                ),

            "ProblemSkill":
                session.scalar(
                    select(
                        func.count()
                    ).select_from(
                        ProblemSkill
                    )
                ),

            "Badge":
                session.scalar(
                    select(
                        func.count()
                    ).select_from(
                        Badge
                    )
                ),
        }

        print(
            "\n=== DATABASE RESULT ==="
        )

        for name, count in counts.items():

            print(
                f"{name}: {count}"
            )

        print(
            "\nSeed completed successfully."
        )

    except Exception:

        session.rollback()
        raise

    finally:

        session.close()


# ============================================================
# CLI
# ============================================================

def main() -> int:

    parser = argparse.ArgumentParser(
        description=(
            "Seed STEPi reference "
            "data from XLSX workbook."
        )
    )

    parser.add_argument(
        "--workbook",
        type=Path,
        default=DEFAULT_WORKBOOK,
    )

    args = parser.parse_args()

    try:

        seed(
            args.workbook
        )

    except Exception as error:

        print(
            "Seed failed; "
            "transaction rolled back: "
            f"{error}",
            file=sys.stderr,
        )

        return 1

    return 0


if __name__ == "__main__":

    raise SystemExit(
        main()
    )